"""
Meridian Packaging International — the Word document with its commercials in prose.

Traps rendered here:
  - every rate is in USD, because the machine-finished liner is imported
  - volume slabs stated as a SENTENCE, and as a RANGE ("12 to 15 per cent"),
    which is not a number a system can silently resolve
  - 90-day payment terms mentioned once, in a closing paragraph
  - tooling for die-cut items quoted separately, INR 18,500 per die, in prose
  - the rate table uses Meridian's own SKU codes, which match nothing
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from contracts.quote import GroundTruth
from tools.labels import meridian as sku, meridian_desc as desc

TEAL = RGBColor(0x0E, 0x4C, 0x4C)


def _p(doc, text, size=9.5, bold=False, italic=False, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold, r.italic = bold, italic
    r.font.name = "Calibri"
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def render(gt: GroundTruth, out_dir: Path) -> Path:
    sub = next(s for s in gt.submissions if s.vendor.vendor_id == "meridian")
    lines = {l.line_no: l for l in gt.rfx.lines}
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Pt(54)

    h = doc.add_paragraph()
    r = h.add_run("MERIDIAN PACKAGING INTERNATIONAL")
    r.bold = True; r.font.size = Pt(15); r.font.color.rgb = TEAL
    _p(doc, "Unit 7, Ambattur Industrial Estate, Chennai 600058, Tamil Nadu, India  |  "
            "GSTIN 33AAGCM8821L1ZK  |  IEC 0488021174", size=7.5, space_after=14)

    _p(doc, "Quotation MPI/EXP/2026/1184", size=10, bold=True, space_after=2)
    _p(doc, "2 September 2026", size=8.5, space_after=12)

    _p(doc, "The Category Purchase Manager\nNandan Consumer Products Ltd\n"
            "Bhiwandi Distribution Centre, Maharashtra", size=9, space_after=12)

    _p(doc, "Dear Sir or Madam,", size=9.5, space_after=8)
    _p(doc, "Sub: Offer against your enquiry " + gt.rfx.rfx_id +
            " for annual corrugated packaging requirement", size=9.5, bold=True, space_after=10)

    # --- the commercials, buried in prose ----------------------------------
    _p(doc,
       "We are pleased to submit our offer against the referenced enquiry. We have "
       "been converting corrugated packaging for the export trade since 1998 and "
       "currently operate two units with a combined converting capacity of 2,200 tonnes "
       "per month.", space_after=8)
    _p(doc,
       "You will observe that our rates are quoted in United States Dollars. This is "
       "deliberate. The machine-finished kraft liner used in our board is imported, and "
       "we are not in a position to carry the currency exposure across an annual "
       "contract. Rates may be converted at the telegraphic transfer selling rate "
       "prevailing on the date of each invoice; for your working, the rate on the date "
       "of this offer is INR 87.10 to the Dollar.", space_after=8)
    _p(doc,
       "The rates set out in the schedule below hold good at an annual offtake above "
       "50,000 pieces per line item. Where the awarded volume on any line falls below "
       "that threshold, an uplift in the region of 12 to 15 per cent would apply to that "
       "line, the precise figure depending on the run length and the changeover involved. "
       "We would of course confirm the applicable figure before despatch.", space_after=8)
    _p(doc,
       "Tooling for the die-cut items is not included in the unit rates. Cutting dies are "
       "charged separately at INR 18,500 per die, payable once against the first order, "
       "and the dies remain your property thereafter.", space_after=12)

    # --- rate schedule ------------------------------------------------------
    _p(doc, "Schedule of rates", size=10, bold=True, space_after=6)
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, htxt in enumerate(["Item code", "Description", "Unit", "Annual qty", "Rate (USD)"]):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(htxt)
        run.bold = True; run.font.size = Pt(8)
    for q in sub.line_quotes:
        l = lines[q.line_no]
        unit = {"piece": "pc", "kg": "kg", "set": "set", "100_pieces": "pc"}[l.uom.value]
        row = t.add_row().cells
        for i, v in enumerate([sku(l), desc(l), unit, f"{l.annual_qty:,}", f"{q.rate:.4f}"]):
            row[i].text = ""
            run = row[i].paragraphs[0].add_run(v)
            run.font.size = Pt(7.5)
            if i == 4:
                row[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()
    # --- closing prose: payment terms appear here and nowhere else ----------
    _p(doc,
       "Rates are ex-works our Chennai unit. Goods and Services Tax is extra as "
       "applicable. Our standard payment terms are 90 days from bill of lading, which "
       "we trust will be acceptable; we are able to discuss an earlier settlement "
       "against an appropriate adjustment. Production lead time is 35 days from receipt "
       "of a firm order, and this offer is valid for 30 days from the date hereof.",
       space_after=8)
    _p(doc,
       "Board strength is certified by bursting strength of 1,400 kPa determined per "
       "TAPPI T810. A test report from our in-house laboratory is enclosed.", space_after=12)

    _p(doc, "We look forward to your favourable consideration.", space_after=14)
    _p(doc, "Yours faithfully,", space_after=18)
    _p(doc, "K. Venkatraman", size=9.5, bold=True, space_after=0)
    _p(doc, "General Manager — International Business", size=8.5, space_after=0)

    out = out_dir / "meridian_offer_MPI-EXP-2026-1184.docx"
    doc.save(out)
    return out
